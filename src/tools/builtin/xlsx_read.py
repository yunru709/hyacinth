"""
name: xlsx_read
description: Read data from an Excel .xlsx file. Returns sheet names, data ranges, and cell values. Supports selecting specific sheets, ranges, and output formats. Requires: pip install openpyxl
parameters:
  type: object
  properties:
    file: {type: string, description: Path to the .xlsx file (absolute or relative to cwd)}
    sheet: {type: string, description: Sheet name to read. If omitted, lists all sheet names and their dimensions.}
    range: {type: string, description: Cell range to read, e.g. A1:D100. If omitted, reads entire used range.}
    max_rows: {type: number, description: Maximum rows to return. Default: 1000. Set to 0 for unlimited.}
    format: {type: string, description: Output format: table (Markdown table, default), csv, or json}
  required: [file]
"""
import json
import os
import sys

def main():
    args = json.loads(sys.stdin.read())
    file_path = args.get('file', '')
    sheet_name = args.get('sheet')
    cell_range = args.get('range')
    max_rows = int(args.get('max_rows', 1000))
    fmt = args.get('format', 'table')

    if not file_path:
        print('Error: file parameter is required.')
        return

    # Resolve path
    if not os.path.isabs(file_path):
        file_path = os.path.join(os.getcwd(), file_path)

    if not os.path.exists(file_path):
        print(f'Error: File not found: {file_path}')
        return

    if not file_path.lower().endswith(('.xlsx', '.xlsm', '.xltx', '.xltm')):
        print(f'Error: File does not appear to be an Excel file: {file_path}')
        return

    try:
        import openpyxl
    except ImportError:
        print('Error: openpyxl is not installed. Run: pip install openpyxl')
        return

    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    except Exception as e:
        print(f'Error: Failed to open workbook: {e}')
        return

    try:
        # No sheet specified → list all sheets with info
        if not sheet_name:
            lines = ['Sheet names:']
            for name in wb.sheetnames:
                ws = wb[name]
                dim = ws.dimensions if ws.dimensions else 'unknown'
                lines.append(f'  - {name}  (range: {dim})')
            print('\n'.join(lines))
            return

        # Validate sheet
        if sheet_name not in wb.sheetnames:
            print(f'Error: Sheet "{sheet_name}" not found. Available: {", ".join(wb.sheetnames)}')
            return

        ws = wb[sheet_name]

        # Determine range
        if cell_range:
            cells = ws[cell_range]
        else:
            cells = ws.iter_rows(values_only=True)

        rows = []
        for row in cells:
            row_values = [cell.value for cell in row] if not isinstance(row, tuple) else list(row)
            rows.append(row_values)
            if max_rows > 0 and len(rows) >= max_rows:
                break

        if not rows:
            print('(empty sheet)')
            return

        if fmt == 'json':
            # Use first row as headers
            headers = [str(h) if h is not None else f'Col{i}' for i, h in enumerate(rows[0])]
            data = []
            for row in rows[1:]:
                obj = {}
                for i, val in enumerate(row):
                    if i < len(headers):
                        obj[headers[i]] = val
                data.append(obj)
            print(json.dumps(data, ensure_ascii=False, indent=2))
        elif fmt == 'csv':
            for row in rows:
                print(','.join(str(c) if c is not None else '' for c in row))
        else:  # table (Markdown)
            if not rows:
                print('(empty sheet)')
                return
            # Header row
            headers = [str(c) if c is not None else '' for c in rows[0]]
            print('| ' + ' | '.join(headers) + ' |')
            print('|' + '|'.join('---' for _ in headers) + '|')
            # Data rows
            for row in rows[1:]:
                vals = [str(c) if c is not None else '' for c in row]
                print('| ' + ' | '.join(vals) + ' |')
            print(f'\n{len(rows)} row(s)')

    finally:
        wb.close()

if __name__ == '__main__':
    main()
